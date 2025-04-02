from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.contrib.auth.models import User
from django.db import transaction
from .models import (
    Empleado, Departamento, RolEmpleado, CriterioEvaluacion,
    Evaluacion, EvaluacionDetalle, Habilidad, Aptitud, Competencia
)
from .forms import (
    UserEmpleadoForm, CriterioEvaluacionForm, EvaluacionForm,
    EvaluacionDetalleForm
)
from django.views.decorators.http import require_http_methods
from rest_framework import viewsets
from rest_framework.permissions import IsAuthenticated
from .serializers import (
    EmpleadoSerializer, DepartamentoSerializer, RolEmpleadoSerializer,
    CriterioEvaluacionSerializer, EvaluacionSerializer
)
from .models import Evaluacion, EvaluacionDetalle, Empleado
from datetime import datetime
from operator import itemgetter
from django.db.models import Avg, Count
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
import simplejson
from django_tables2 import RequestConfig, SingleTableView
from django_tables2.export.views import ExportMixin
from django_filters.views import FilterView
from .tables import UserEmpleadoTable, CriterioEvaluacionTable, EvaluacionTable
from .filters import UserEmpleadoFilter, CriterioEvaluacionFilter, EvaluacionFilter
import uuid

# Existing view functions...

class FilteredTableView(ExportMixin, FilterView, SingleTableView):
    """
    Vista genérica para tablas filtradas con capacidad de exportación
    """
    template_name = 'empleados/filtered_table.html'
    export_formats = ['csv', 'xlsx', 'json']

# ViewSets
class EmpleadoViewSet(viewsets.ModelViewSet):
    queryset = Empleado.objects.all()
    serializer_class = EmpleadoSerializer
    permission_classes = [IsAuthenticated]

class DepartamentoViewSet(viewsets.ModelViewSet):
    queryset = Departamento.objects.all()
    serializer_class = DepartamentoSerializer
    permission_classes = [IsAuthenticated]

class RolEmpleadoViewSet(viewsets.ModelViewSet):
    queryset = RolEmpleado.objects.all()
    serializer_class = RolEmpleadoSerializer
    permission_classes = [IsAuthenticated]

class CriterioEvaluacionViewSet(viewsets.ModelViewSet):
    queryset = CriterioEvaluacion.objects.all()
    serializer_class = CriterioEvaluacionSerializer
    permission_classes = [IsAuthenticated]

class EvaluacionViewSet(viewsets.ModelViewSet):
    queryset = Evaluacion.objects.all()
    serializer_class = EvaluacionSerializer
    permission_classes = [IsAuthenticated]

# Keep all existing view functions...

@login_required
def gestionar_trabajadores(request):
    # Filtro para usuarios
    filterset = UserEmpleadoFilter(request.GET, queryset=User.objects.filter(is_active=True))
    
    # Tabla con los resultados filtrados
    table = UserEmpleadoTable(filterset.qs)
    RequestConfig(request, paginate={"per_page": 10}).configure(table)
    
    # Datos adicionales para el formulario
    departamentos = Departamento.objects.all()
    roles = RolEmpleado.objects.all()
    habilidades = Habilidad.objects.all()
    aptitudes = Aptitud.objects.all()
    competencias = Competencia.objects.all()
    
    return render(request, 'empleados/gestionar_trabajadores.html', {
        'table': table,
        'filter': filterset,
        'usuarios': filterset.qs,
        'departamentos': departamentos,
        'roles': roles,
        'habilidades': habilidades,
        'aptitudes': aptitudes,
        'competencias': competencias
    })

@login_required
@require_http_methods(["POST"])
def crear_editar_trabajador(request):
    user_id = request.POST.get('user_id')
    if user_id:
        user = get_object_or_404(User, id=user_id)
    else:
        user = None
    
    if request.method == 'POST':
        # Si es un nuevo usuario, generar un nombre de usuario único
        if not user:
            # Generar un nombre de usuario basado en el nombre y apellido o un UUID si no están disponibles
            first_name = request.POST.get('first_name', '')
            last_name = request.POST.get('last_name', '')
            if first_name and last_name:
                username = f"{first_name.lower()[0]}{last_name.lower().replace(' ', '')}"
                # Verificar si el nombre de usuario ya existe
                if User.objects.filter(username=username).exists():
                    username = f"{username}{User.objects.count()}"
            else:
                # Si no hay nombre o apellido, usar un UUID
                username = str(uuid.uuid4())[:8]
            
            # Crear un nuevo usuario con el nombre de usuario generado
            user = User(username=username)
        
        # Actualizar los campos del usuario
        user.first_name = request.POST.get('first_name', '')
        user.last_name = request.POST.get('last_name', '')
        user.email = request.POST.get('email', '')
        user.is_active = 'activo' in request.POST
        user.save()
        
        # Actualizar o crear el empleado asociado
        empleado, created = Empleado.objects.get_or_create(user=user)
        
        # Actualizar departamento y rol
        departamento_id = request.POST.get('departamento')
        if departamento_id:
            empleado.departamento = Departamento.objects.get(id=departamento_id)
        else:
            empleado.departamento = None
            
        rol_id = request.POST.get('rol')
        if rol_id:
            empleado.rol = RolEmpleado.objects.get(id=rol_id)
        else:
            empleado.rol = None
            
        empleado.save()
        
        # Actualizar habilidades, aptitudes y competencias
        habilidades_ids = request.POST.getlist('habilidades')
        aptitudes_ids = request.POST.getlist('aptitudes')
        competencias_ids = request.POST.getlist('competencias')
        
        empleado.habilidades.clear()
        if habilidades_ids:
            empleado.habilidades.add(*Habilidad.objects.filter(id__in=habilidades_ids))
            
        empleado.aptitudes.clear()
        if aptitudes_ids:
            empleado.aptitudes.add(*Aptitud.objects.filter(id__in=aptitudes_ids))
            
        empleado.competencias.clear()
        if competencias_ids:
            empleado.competencias.add(*Competencia.objects.filter(id__in=competencias_ids))
        
        return JsonResponse({'status': 'success', 'message': 'Usuario guardado exitosamente.'})
    
    return JsonResponse({'status': 'error', 'message': 'Método no permitido'}, status=405)

@login_required
@require_http_methods(["POST"])
def eliminar_trabajador(request, user_id):
    user = get_object_or_404(User, id=user_id)
    user.is_active = False
    user.save()
    return JsonResponse({'status': 'success', 'message': 'Usuario eliminado exitosamente.'})

@login_required
@require_http_methods(["GET"])
def obtener_trabajador(request, user_id):
    user = get_object_or_404(User, id=user_id)
    try:
        empleado = user.empleado
        habilidades = list(empleado.habilidades.values_list('id', flat=True))
        aptitudes = list(empleado.aptitudes.values_list('id', flat=True))
        competencias = list(empleado.competencias.values_list('id', flat=True))
        departamento = empleado.departamento.id if empleado.departamento else None
        rol = empleado.rol.id if empleado.rol else None
    except Empleado.DoesNotExist:
        habilidades = []
        aptitudes = []
        competencias = []
        departamento = None
        rol = None
    
    data = {
        'id': user.id,
        'username': user.username,
        'first_name': user.first_name,
        'last_name': user.last_name,
        'email': user.email,
        'activo': user.is_active,
        'departamento': departamento,
        'rol': rol,
        'habilidades': habilidades,
        'aptitudes': aptitudes,
        'competencias': competencias
    }
    
    return JsonResponse(data)

@login_required
def gestionar_criterios_evaluacion(request):
    # Filtro para criterios
    filterset = CriterioEvaluacionFilter(request.GET, queryset=CriterioEvaluacion.objects.all())
    
    # Tabla con los resultados filtrados
    table = CriterioEvaluacionTable(filterset.qs)
    RequestConfig(request, paginate={"per_page": 10}).configure(table)
    
    # Datos adicionales para el formulario
    roles = RolEmpleado.objects.all()
    
    return render(request, 'empleados/gestionar_criterios_evaluacion.html', {
        'table': table,
        'filter': filterset,
        'criterios': filterset.qs,
        'roles': roles
    })

@login_required
@require_http_methods(["POST"])
def crear_editar_criterio(request):
    criterio_id = request.POST.get('criterio_id')
    if criterio_id:
        criterio = get_object_or_404(CriterioEvaluacion, id=criterio_id)
        form = CriterioEvaluacionForm(request.POST, instance=criterio)
    else:
        form = CriterioEvaluacionForm(request.POST)

    if form.is_valid():
        form.save()
        return JsonResponse({'status': 'success'})
    else:
        return JsonResponse({'status': 'error', 'errors': form.errors})

@login_required
@require_http_methods(["POST"])
def eliminar_criterio(request, criterio_id):
    criterio = get_object_or_404(CriterioEvaluacion, id=criterio_id)
    criterio.delete()
    return JsonResponse({'status': 'success'})

@login_required
@require_http_methods(["GET"])
def obtener_criterio(request, criterio_id):
    criterio = get_object_or_404(CriterioEvaluacion, id=criterio_id)
    data = {
        'id': criterio.id,
        'denominacion': criterio.denominacion,
        'rol': criterio.rol.id,
        'generico': criterio.generico,
    }
    return JsonResponse(data)

@login_required
def gestionar_evaluaciones(request):
    # Filtro para evaluaciones
    filterset = EvaluacionFilter(request.GET, queryset=Evaluacion.objects.all())
    
    # Tabla con los resultados filtrados
    table = EvaluacionTable(filterset.qs)
    RequestConfig(request, paginate={"per_page": 10}).configure(table)
    
    return render(request, 'empleados/gestionar_evaluaciones.html', {
        'table': table,
        'filter': filterset,
        'evaluaciones': filterset.qs
    })

@login_required
def crear_evaluacion(request):
    if request.method == 'POST':
        form = EvaluacionForm(request.POST)
        if form.is_valid():
            evaluacion = form.save(commit=False)
            evaluacion.autor = request.user
            evaluacion.save()
            return redirect('editar_evaluacion', evaluacion_id=evaluacion.id)
    else:
        form = EvaluacionForm()
    
    return render(request, 'empleados/crear_evaluacion.html', {
        'form': form
    })

@login_required
def editar_evaluacion(request, evaluacion_id):
    evaluacion = get_object_or_404(Evaluacion, id=evaluacion_id)
    empleados = Empleado.objects.filter(user__activo=True)
    criterios = CriterioEvaluacion.objects.all()
    
    if request.method == 'POST':
        with transaction.atomic():
            for empleado in empleados:
                for criterio in criterios:
                    puntuacion = request.POST.get(f'puntuacion_{empleado.id}_{criterio.id}')
                    
                    if puntuacion:
                        EvaluacionDetalle.objects.update_or_create(
                            evaluacion=evaluacion,
                            empleado=empleado,
                            criterio=criterio,
                            defaults={
                                'puntuacion': int(puntuacion),
                            }
                        )
            
            return redirect('gestionar_evaluaciones')
    
    # Obtener evaluaciones existentes
    evaluaciones_data = {}
    for detalle in evaluacion.detalles.all():
        if detalle.empleado.id not in evaluaciones_data:
            evaluaciones_data[detalle.empleado.id] = {}
        evaluaciones_data[detalle.empleado.id][detalle.criterio.id] = detalle.puntuacion
    
    # Serializar los datos usando simplejson
    evaluaciones_json = simplejson.dumps(evaluaciones_data)
    
    print("Loaded evaluaciones:", evaluaciones_json)  # Debug: Verificar evaluaciones cargadas
    
    return render(request, 'empleados/editar_evaluacion.html', {
        'evaluacion': evaluacion,
        'empleados': empleados,
        'criterios': criterios,
        'evaluaciones_json': evaluaciones_json
    })

@login_required
def eliminar_evaluacion(request, evaluacion_id):
    evaluacion = get_object_or_404(Evaluacion, id=evaluacion_id)
    if request.method == 'POST':
        evaluacion.delete()
        return JsonResponse({'status': 'success'})
    return JsonResponse({'status': 'error'}, status=400)

@login_required
def calculo_puntuaciones(request):
    if request.method == 'POST':
        fecha_inicial = request.POST.get('fecha_inicial')
        fecha_final = request.POST.get('fecha_final')
        
        if not fecha_inicial or not fecha_final:
            return JsonResponse({'error': 'Fechas no proporcionadas'}, status=400)

        fecha_inicial = datetime.strptime(fecha_inicial, '%Y-%m-%d').date()
        fecha_final = datetime.strptime(fecha_final, '%Y-%m-%d').date()

        evaluaciones = Evaluacion.objects.filter(fecha__range=(fecha_inicial, fecha_final))
        
        if 'evaluaciones' in request.POST:
            evaluaciones_ids = request.POST.getlist('evaluaciones')
            evaluaciones = evaluaciones.filter(id__in=evaluaciones_ids)
            
            resultados = []
            for empleado in Empleado.objects.filter(evaluaciondetalle__evaluacion__in=evaluaciones).distinct():
                detalles = EvaluacionDetalle.objects.filter(
                    evaluacion__in=evaluaciones,
                    empleado=empleado
                )
                
                # Calcular promedio incluyendo todos los criterios
                promedio_total = detalles.aggregate(Avg('puntuacion'))['puntuacion__avg']
                
                # Calcular promedio excluyendo criterios genéricos
                promedio_no_generico = detalles.exclude(criterio__generico=True).aggregate(Avg('puntuacion'))['puntuacion__avg']
                
                resultados.append({
                    'nombre': empleado.user.get_full_name(),
                    'promedio': round(promedio_total, 2) if promedio_total else 0,
                    'promedio_no_generico': round(promedio_no_generico, 2) if promedio_no_generico else 0,
                })
            
            # Ordenar resultados por promedio (todos los criterios) de forma descendente
            resultados.sort(key=lambda x: x['promedio'], reverse=True)
            
            if 'export' in request.POST:
                return export_to_excel(resultados, 'calculo_puntuaciones.xlsx')
            
            return JsonResponse({'empleados': resultados})
        else:
            # Búsqueda de evaluaciones
            evaluaciones_data = [{
                'id': eval.id,
                'fecha': eval.fecha.strftime('%Y-%m-%d'),
                'mes_inicial': eval.mes_inicial.strftime('%Y-%m-%d'),
                'mes_final': eval.mes_final.strftime('%Y-%m-%d')
            } for eval in evaluaciones]
            return JsonResponse({'evaluaciones': evaluaciones_data})

    return render(request, 'empleados/calculo_puntuaciones.html')

def export_to_excel(data, filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resultados"

    # Escribir encabezados
    headers = ['Empleado', 'Promedio (todos los criterios)', 'Promedio (sin criterios genéricos)']
    for col_num, header in enumerate(headers, 1):
        col_letter = get_column_letter(col_num)
        ws[f'{col_letter}1'] = header
        ws[f'{col_letter}1'].font = Font(bold=True)

    # Escribir datos
    for row_num, empleado in enumerate(data, 2):
        ws.cell(row=row_num, column=1, value=empleado['nombre'])
        ws.cell(row=row_num, column=2, value=empleado['promedio'])
        ws.cell(row=row_num, column=3, value=empleado['promedio_no_generico'])

    # Ajustar ancho de columnas
    for column_cells in ws.columns:
        length = max(len(str(cell.value)) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = length + 2

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename={filename}'
    wb.save(response)
    return response

@login_required
def reportes_analisis(request):
    if request.method == 'POST':
        fecha_inicial = request.POST.get('fecha_inicial')
        fecha_final = request.POST.get('fecha_final')
        
        if not fecha_inicial or not fecha_final:
            return JsonResponse({'error': 'Fechas no proporcionadas'}, status=400)

        fecha_inicial = datetime.strptime(fecha_inicial, '%Y-%m-%d').date()
        fecha_final = datetime.strptime(fecha_final, '%Y-%m-%d').date()

        evaluaciones = Evaluacion.objects.filter(fecha__range=(fecha_inicial, fecha_final))
        
        if 'evaluaciones' in request.POST:
            evaluaciones_ids = request.POST.getlist('evaluaciones')
            evaluaciones = evaluaciones.filter(id__in=evaluaciones_ids)
            
            # Datos para la gráfica de barras
            empleados_por_puntuacion = EvaluacionDetalle.objects.filter(
                evaluacion__in=evaluaciones
            ).values('puntuacion').annotate(
                cantidad=Count('empleado', distinct=True)
            ).order_by('puntuacion')

            bar_labels = [f"Puntuación {p['puntuacion']}" for p in empleados_por_puntuacion]
            bar_data = [p['cantidad'] for p in empleados_por_puntuacion]

            # Datos para la gráfica de pastel
            criterios_evaluaciones = EvaluacionDetalle.objects.filter(
                evaluacion__in=evaluaciones
            ).values('criterio__denominacion', 'puntuacion').annotate(
                cantidad=Count('id')
            ).order_by('criterio__denominacion', 'puntuacion')

            pie_data = {}
            for ce in criterios_evaluaciones:
                criterio = ce['criterio__denominacion']
                puntuacion = ce['puntuacion']
                cantidad = ce['cantidad']
                if criterio not in pie_data:
                    pie_data[criterio] = {'labels': [], 'data': []}
                pie_data[criterio]['labels'].append(f"Puntuación {puntuacion}")
                pie_data[criterio]['data'].append(cantidad)

            if 'export' in request.POST:
                return export_to_excel_analisis(empleados_por_puntuacion, criterios_evaluaciones, 'reportes_analisis.xlsx')
            
            return JsonResponse({
                'bar_chart': {'labels': bar_labels, 'data': bar_data},
                'pie_charts': pie_data
            })
        else:
            # Búsqueda de evaluaciones
            evaluaciones_data = [{
                'id': eval.id,
                'fecha': eval.fecha.strftime('%Y-%m-%d'),
                'mes_inicial': eval.mes_inicial.strftime('%Y-%m-%d'),
                'mes_final': eval.mes_final.strftime('%Y-%m-%d')
            } for eval in evaluaciones]
            return JsonResponse({'evaluaciones': evaluaciones_data})

    return render(request, 'empleados/reportes_analisis.html')

def export_to_excel_analisis(empleados_por_puntuacion, criterios_evaluaciones, filename):
    wb = openpyxl.Workbook()
    
    # Hoja para empleados por puntuación
    ws_empleados = wb.active
    ws_empleados.title = "Empleados por Puntuación"
    ws_empleados.append(["Puntuación", "Cantidad de Empleados"])
    for item in empleados_por_puntuacion:
        ws_empleados.append([item['puntuacion'], item['cantidad']])

    # Hoja para criterios y evaluaciones
    ws_criterios = wb.create_sheet("Criterios y Evaluaciones")
    ws_criterios.append(["Criterio", "Puntuación", "Cantidad"])
    for item in criterios_evaluaciones:
        ws_criterios.append([item['criterio__denominacion'], item['puntuacion'], item['cantidad']])

    # Ajustar ancho de columnas
    for sheet in [ws_empleados, ws_criterios]:
        for column_cells in sheet.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = length + 2

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename={filename}'
    wb.save(response)
    return response
